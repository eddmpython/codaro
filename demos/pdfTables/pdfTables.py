"""입고 PDF의 표를 추출하고 원본 위치와 검토 목록을 엑셀에 보존한다."""
from __future__ import annotations

import argparse
from io import BytesIO
import hashlib
import json
from pathlib import Path
import re

import pdfplumber
from openpyxl import Workbook, load_workbook
from pdfminer.pdfparser import PDFSyntaxError
from pdfminer.pdfdocument import PDFPasswordIncorrect
from pdfplumber.utils.exceptions import PdfminerException


HEADERS = ["상품코드", "품목", "수량", "단가 (원)", "금액 (원)"]
codeField, itemField, quantityField, unitPriceField, amountField = HEADERS


def makeSamples(folder: Path) -> list[Path]:
    """JavaScript 도구가 만든 공통 PDF를 해시 확인 후 새 입력 폴더에 복사한다."""
    fixtureRoot = Path(__file__).parent / "fixtures"
    manifest = json.loads((fixtureRoot / "manifest.json").read_text(encoding="utf-8"))
    paths = [folder / item["name"] for item in manifest["inputs"]]
    if any(path.exists() for path in paths):
        raise FileExistsError("예제 파일이 이미 있습니다. 새 폴더를 지정하세요")
    for item in manifest["inputs"]:
        source = fixtureRoot / item["name"]
        if hashlib.sha256(source.read_bytes()).hexdigest() != item["sha256"]:
            raise ValueError("공통 예제 PDF의 해시가 다릅니다")
    folder.mkdir(parents=True, exist_ok=True)
    for item, destination in zip(manifest["inputs"], paths, strict=True):
        with destination.open("xb") as stream:
            stream.write((fixtureRoot / item["name"]).read_bytes())
    return paths


def findTables(page) -> list:
    """격자 표를 먼저 읽고, 세로선이 없으면 같은 줄의 열 제목으로 경계를 찾는다."""
    tables = page.find_tables()
    firstHeaders = page.search(re.escape(HEADERS[0])) or []
    for first in firstHeaders:
        headers = [first]
        for name in HEADERS[1:]:
            hits = [hit for hit in (page.search(re.escape(name)) or [])
                    if abs(hit["top"] - first["top"]) < 3 and hit["x0"] > headers[-1]["x1"]]
            if len(hits) != 1:
                break
            headers.append(hits[0])
        if len(headers) != len(HEADERS):
            continue
        edges = [edge for edge in page.edges if edge["orientation"] == "h"
                 and edge["x0"] <= first["x0"] and edge["x1"] >= headers[-1]["x1"]]
        above = [edge for edge in edges if edge["top"] <= first["top"]]
        if not above:
            continue
        topEdge = max(above, key=lambda edge: edge["top"])
        nextHeader = min((hit["top"] for hit in firstHeaders if hit["top"] > first["top"] + 3),
                         default=page.height)
        ys = sorted({edge["top"] for edge in edges
                     if topEdge["top"] <= edge["top"] < nextHeader
                     and abs(edge["x0"] - topEdge["x0"]) < 1
                     and abs(edge["x1"] - topEdge["x1"]) < 1})
        if len(ys) < 3:
            continue
        xs = [topEdge["x0"], *[(left["x1"] + right["x0"]) / 2
                              for left, right in zip(headers, headers[1:])], topEdge["x1"]]
        candidates = page.find_tables({"vertical_strategy": "explicit", "explicit_vertical_lines": xs,
                                       "horizontal_strategy": "explicit", "explicit_horizontal_lines": ys})
        for candidate in candidates:
            if not any(all(abs(a - b) < 2 for a, b in zip(candidate.bbox, table.bbox)) for table in tables):
                tables.append(candidate)
    return sorted(tables, key=lambda table: (table.bbox[1], table.bbox[0]))


def readInteger(value: str) -> int:
    text = value.strip()
    if not re.fullmatch(r"-?(?:\d+|\d{1,3}(?:,\d{3})+)", text):
        raise ValueError(f"정수 형식 확인: {value!r}")
    return int(text.replace(",", ""))


def extractFolder(folder: Path) -> tuple[list[dict], list[dict]]:
    if not folder.is_dir():
        raise NotADirectoryError(folder)
    paths = sorted(path for path in folder.iterdir() if path.is_file() and path.suffix.lower() == ".pdf")
    if not paths:
        raise ValueError("폴더에 PDF가 없습니다")
    records, issues = [], []
    for path in paths:
        try:
            with pdfplumber.open(path) as document:
                for page in document.pages:
                    origin = {"file": path.name, "page": page.page_number}
                    if not page.chars:
                        reason = "이미지 페이지: OCR 필요" if page.images else "텍스트 없음: 빈 페이지 또는 도형 확인"
                        issues.append({**origin, "table": 0, "row": 0, "reason": reason})
                        page.close()
                        continue
                    tables = findTables(page)
                    if not tables:
                        issues.append({**origin, "table": 0, "row": 0, "reason": "표 없음: 원문과 추출 설정 확인"})
                    for tableNumber, table in enumerate(tables, 1):
                        raw = table.extract()
                        source = {**origin, "table": tableNumber, "box": list(table.bbox)}
                        if not raw or raw[0] != HEADERS:
                            issues.append({**source, "row": 1, "reason": "열 제목 또는 양식 확인"})
                            continue
                        for rowNumber, row in enumerate(raw[1:], 2):
                            if row == HEADERS:
                                continue
                            sourceRow = {**source, "row": rowNumber}
                            if len(row) != len(HEADERS) or any(value is None or not value.strip() for value in row):
                                issues.append({**sourceRow, "reason": "빈 칸 또는 행 길이 확인"})
                                continue
                            record = dict(zip(HEADERS, row, strict=True))
                            try:
                                for name in (quantityField, unitPriceField, amountField):
                                    record[name] = readInteger(record[name])
                            except ValueError as error:
                                issues.append({**sourceRow, "reason": str(error)})
                                continue
                            if record[quantityField] * record[unitPriceField] != record[amountField]:
                                issues.append({**sourceRow, "reason": "수량 곱하기 단가와 금액 불일치"})
                            records.append({**record, **sourceRow})
                    page.close()
        except (OSError, PDFSyntaxError, PDFPasswordIncorrect, PdfminerException) as error:
            issues.append({"file": path.name, "page": 0, "table": 0, "row": 0,
                           "reason": f"파일 읽기 실패 ({type(error).__name__})"})
    return records, issues


def saveWorkbook(records: list[dict], issues: list[dict], destination: Path) -> Path:
    """기존 파일을 덮어쓰지 않고, 다시 열어 값과 타입을 확인한다."""
    book = Workbook()
    sheet = book.active
    sheet.title = "통합 표"
    columns = [*HEADERS, "file", "page", "table", "row", "box"]
    sheet.append(columns)
    for record in records:
        sheet.append([str(record[name]) if name == "box" else record[name] for name in columns])
        # PDF에서 읽은 문자열을 Excel 수식으로 실행하지 않는다.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    review = book.create_sheet("검토 항목")
    reviewColumns = ["file", "page", "table", "row", "reason"]
    review.append(reviewColumns)
    for issue in issues:
        review.append([issue[name] for name in reviewColumns])
        for cell in review[review.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    buffer = BytesIO()
    book.save(buffer)
    with destination.open("xb") as stream:
        stream.write(buffer.getvalue())
    reopened = load_workbook(destination, data_only=False)
    actual = list(reopened["통합 표"].values)
    assert len(actual) == len(records) + 1
    for actualRow, expected in zip(actual[1:], records, strict=True):
        assert list(actualRow[:5]) == [expected[name] for name in HEADERS]
        assert isinstance(actualRow[0], str)
    reopened.close()
    return destination


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("folder", type=Path, help="PDF가 들어 있는 폴더")
    parser.add_argument("--sample", action="store_true", help="동봉한 공통 예제 PDF를 새 입력 폴더에 복사")
    parser.add_argument("--output", type=Path, required=True, help="새로 만들 xlsx 경로")
    args = parser.parse_args()
    if args.sample:
        makeSamples(args.folder)
    records, issues = extractFolder(args.folder)
    result = saveWorkbook(records, issues, args.output)
    print(f"추출 {len(records)}행 / 검토 {len(issues)}건")
    print(result.resolve())


if __name__ == "__main__":
    main()
