from __future__ import annotations

import contextlib
import importlib.util
from io import StringIO
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest

import yaml
from openpyxl import load_workbook
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[2]
LESSON = ROOT / "curricula/python/automation/office/pdf/04_표추출.yaml"
SPEC = importlib.util.spec_from_file_location("pdfTablesPractice", ROOT / "demos/pdfTables/pdfTables.py")
assert SPEC and SPEC.loader
DEMO = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(DEMO)


class PdfBlogPracticeTest(unittest.TestCase):
    def testGridTableDoesNotHideOtherTableOnSamePage(self):
        from types import SimpleNamespace
        import pdfplumber
        with pdfplumber.open(ROOT / "demos/pdfTables/fixtures/입고내역_창고A.pdf") as document:
            page = document.pages[0]

            class MixedPage:
                def __getattr__(self, name):
                    return getattr(page, name)

                def find_tables(self, settings=None):
                    if settings is None:
                        return [SimpleNamespace(bbox=(10, 10, 20, 20))]
                    return page.find_tables(settings)

            tables = DEMO.findTables(MixedPage())
            self.assertEqual(len(tables), 2)
            self.assertEqual(tables[1].extract()[1][0], "0012")

    def testLessonExamplesAndSolutionsExecute(self):
        lesson = yaml.safe_load(LESSON.read_text(encoding="utf-8"))
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            DEMO.makeSamples(folder / "input")
            shutil.copyfile(ROOT / "demos/pdfTables/pdfTables.py", folder / "pdfTables.py")
            previous = Path.cwd()
            sys.path.insert(0, str(folder))
            try:
                os.chdir(folder)
                for section in lesson["sections"]:
                    for code in [section.get("snippet"), section.get("exercise", {}).get("solution")]:
                        if code:
                            with self.subTest(section=section["id"]), contextlib.redirect_stdout(StringIO()):
                                exec(compile(code, section["id"], "exec"), {})
            finally:
                os.chdir(previous)
                sys.path.remove(str(folder))
                sys.modules.pop("pdfTables", None)

    def testAllFilesPagesAndWorkbookRoundTrip(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            DEMO.makeSamples(folder)
            records, issues = DEMO.extractFolder(folder)
            fixture = json.loads((ROOT / "demos/pdfTables/fixtures/manifest.json").read_text(encoding="utf-8"))
            expected = [dict(zip(DEMO.HEADERS, values, strict=True), file=table["file"], page=table["page"], table=table["table"])
                        for table in fixture["tables"] for values in table["rows"]]
            self.assertEqual([{key: row[key] for key in expected[0]} for row in records], expected)
            self.assertEqual(issues, [])
            self.assertEqual({(row["file"], row["page"]) for row in records},
                             {(table["file"], table["page"]) for table in fixture["tables"]})
            self.assertEqual(sum(row[DEMO.amountField] for row in records),
                             sum(row[DEMO.amountField] for row in expected))
            output = DEMO.saveWorkbook(records, issues, folder / "result.xlsx")
            book = load_workbook(output)
            self.assertEqual(book["통합 표"]["A2"].value, "0012")
            self.assertEqual(book["통합 표"]["E2"].value, 456000)
            self.assertEqual(book["검토 항목"].max_row, 1)
            book.close()
            with self.assertRaises(FileExistsError):
                DEMO.saveWorkbook(records, issues, output)
            with self.assertRaises(FileExistsError):
                DEMO.makeSamples(folder)

    def testImageEmptyMalformedAndUnmatchedPagesRemainVisible(self):
        from PIL import Image
        from reportlab.lib.utils import ImageReader
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            canvas = Canvas(str(folder / "scan.PDF"))
            canvas.drawImage(ImageReader(Image.new("RGB", (50, 50), "white")), 60, 600)
            canvas.showPage()
            canvas.showPage()
            canvas.drawString(60, 600, "Text without a table")
            canvas.save()
            (folder / "broken.pdf").write_bytes(b"not a PDF")
            records, issues = DEMO.extractFolder(folder)
            self.assertEqual(records, [])
            reasons = [issue["reason"] for issue in issues]
            self.assertTrue(any("OCR" in reason for reason in reasons))
            self.assertTrue(any("빈 페이지" in reason for reason in reasons))
            self.assertTrue(any("표 없음" in reason for reason in reasons))
            self.assertTrue(any("파일 읽기 실패" in reason for reason in reasons))

    def testNumbersAndFormulaText(self):
        self.assertEqual(DEMO.readInteger("1,234"), 1234)
        for value in ["1,2,3", "=1+1", "", "1.5", "1O0"]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                DEMO.readInteger(value)
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            DEMO.makeSamples(folder)
            records, issues = DEMO.extractFolder(folder)
            records[0][DEMO.itemField] = "=1+1"
            bookPath = DEMO.saveWorkbook(records, issues, folder / "safe.xlsx")
            book = load_workbook(bookPath, data_only=False)
            self.assertEqual(book["통합 표"]["B2"].data_type, "s")
            book.close()

    def testSectionImageReferencesAndLocalBytes(self):
        import hashlib
        lesson = yaml.safe_load(LESSON.read_text(encoding="utf-8"))
        manifest = json.loads((ROOT / "assets/brand/visuals/manifest.json").read_text(encoding="utf-8"))
        assets = {asset["id"]: asset for asset in manifest["assets"]}
        ids = []
        for section in lesson["sections"]:
            images = [block for block in section.get("blocks", []) if block["type"] == "image"]
            if section["id"] not in {"prepare", "references"}:
                self.assertEqual(len(images), 1, section["id"])
            for block in images:
                self.assertEqual(block["placement"], "sectionLead")
                asset = assets[block["assetId"]]
                self.assertEqual(asset["sourceType"], "sharedRaster")
                self.assertEqual(asset["learning"]["lessonRefs"], ["pdf/04_표추출"])
                actual = hashlib.sha256((ROOT / asset["sourcePath"]).read_bytes()).hexdigest()
                self.assertEqual(asset["sourceHash"], "sha256-" + actual)
                ids.append(block["assetId"])
        self.assertEqual(len(ids), len(set(ids)))


if __name__ == "__main__":
    unittest.main()
